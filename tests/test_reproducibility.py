"""Checks that catch incorrect recovery statistics and accidental data leakage."""
from pathlib import Path
import json,sys,tempfile,unittest
ROOT=Path(__file__).resolve().parents[1]
sys.dont_write_bytecode=True
sys.path.insert(0,str(ROOT/'code'))
import numpy as np
import pandas as pd
import descriptors as D
import ml_v3 as M
import retrospective_bo as R
import recovery


class ReproducibilityTests(unittest.TestCase):
    def test_archive_and_censored_median(self):
        runs=pd.read_csv(ROOT/'data/reference/retrospective/runs.csv')
        pool=R.load_pool()
        self.assertEqual(recovery.validate_runs(runs,pool[M.COL_X]),100)
        curve,summary=recovery.tabulate(runs)
        for strategy,n,value in [('GP–EHVI',10,57),('GP–EHVI',17,100),('Grid scan',20,100),('Random',30,36)]:
            actual=curve.loc[(curve.Strategy==strategy)&(curve.Queries==n),'Recovery_percent'].iloc[0]
            self.assertAlmostEqual(actual,value)
        self.assertTrue(pd.isna(summary.loc[summary.Strategy=='Random','Median_queries_all_runs'].iloc[0]))
        self.assertTrue(np.isnan(recovery.cohort_median([1,2,np.nan,np.nan])))
        self.assertEqual(recovery.cohort_median([1,2,3,np.nan]),2.5)

    def test_reject_mismatched_initial_set(self):
        runs=pd.read_csv(ROOT/'data/reference/retrospective/runs.csv')
        runs.loc[0,'Start_x']='0.050;0.052;0.054'
        with self.assertRaises(ValueError):recovery.validate_runs(runs,R.load_pool()[M.COL_X])

    def test_acquisition_does_not_read_unqueried_outcomes(self):
        pool=R.load_pool();known=[0,20,80]
        changed=pool.copy();mask=~changed.index.isin(known)
        for col in [M.COL_EB,M.COL_WE,M.COL_WREC,M.COL_ETA]:changed.loc[mask,col]=-12345.0
        prior=R.MC_SAMPLES;R.MC_SAMPLES=64
        try:
            first=R.ehvi_next(pool,known,np.random.default_rng(2042))
            second=R.ehvi_next(changed,known,np.random.default_rng(2042))
            self.assertEqual(first,second)
            self.assertNotIn(first,known)
        finally:R.MC_SAMPLES=prior

    def test_descriptor_missing_value_fails(self):
        from openpyxl import load_workbook
        book=load_workbook(M.DATA_XLSX)
        sheet=book['ML描述符'];headers=[c.value for c in sheet[1]]
        column=headers.index(D.DESCRIPTOR_COLS[0])+1
        sheet.cell(row=2,column=column,value='invalid')
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'invalid.xlsx';book.save(path)
            with self.assertRaises(ValueError):D.validate(path)


if __name__=='__main__':unittest.main()
